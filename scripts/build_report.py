"""Informe editorial Word/PDF desde fuentes del repositorio, sin aprobar gates."""
import argparse
import hashlib
import json
import re
import subprocess
from html import escape
from pathlib import Path

import yaml
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, Frame, Image, KeepTogether, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.platypus.tableofcontents import TableOfContents

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables/revision-editorial'
REPO = 'https://github.com/josemanuelcondeacosta-lang/nexoentrega-360-sdd-ia'
GREEN = '193E35'
blocks = []


def add(kind, value=None):
    blocks.append((kind, value))


def heading(text, level=1):
    add('h' + str(level), text)


def paragraph(text):
    add('p', text)


def table(headers, rows):
    add('table', [headers] + rows)


def source(path, sha):
    add('link', (path, f'{REPO}/blob/{sha}/{path}'))


def markdown(path, min_level=2):
    """Preserva contenido; convierte Markdown a estilos y tablas reales."""
    lines = (ROOT / path).read_text(encoding='utf-8-sig').splitlines()
    in_code = False
    pending = []
    for line in lines + ['']:
        if line.startswith('|'):
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+', c) for c in cells):
                pending.append(cells)
            continue
        if pending:
            if len({len(row) for row in pending}) == 1:
                add('table', pending)
            else:
                for row in pending:
                    paragraph(' · '.join(row))
            pending = []
        if line.startswith('```'):
            in_code = not in_code
            continue
        if not line.strip():
            continue
        if in_code:
            add('code', line)
        elif line.startswith('#'):
            n = len(line) - len(line.lstrip('#'))
            heading(line.lstrip('# '), min(3, min_level + n - 1))
        elif re.match(r'^[-*] ', line):
            add('bullet', line[2:])
        else:
            paragraph(line.removeprefix('> '))


def content(sha):
    github = json.loads((ROOT / 'docs/evidence/github/snapshot.json').read_text(encoding='utf-8'))
    current_ci = json.loads((ROOT / 'docs/evidence/github/current-ci.json').read_text(encoding='utf-8'))
    add('cover', sha)
    add('toc')
    heading('1. Entrega identificada y estado del proyecto')
    paragraph('NexoEntrega 360 es un sistema académico para coordinar la distribución comunitaria de kits en dos organizaciones de Barranquilla. Este informe reúne especificaciones, arquitectura, contratos, uso de IA, resultados y participación, siguiendo los nueve apartados del anexo oficial 09.')
    add('callout', 'VERSIÓN PROVISIONAL PARA REVISIÓN. No se declara aprobado G1/G2 ni aceptada la entrega. Las decisiones humanas pendientes y los requisitos sin evidencia se identifican expresamente.')
    table(['Elemento', 'Identificación'], [
        ['Repositorio', REPO], ['Código de esta edición', sha],
        ['Base de pruebas anteriores', 'f33183aa72c261bb513cf8aada1113605b461262'],
        ['Ejecución local', 'http://127.0.0.1:8000 — requiere Docker; no es URL pública para el docente'],
        ['release-examen / baseline-v1', 'No creados como aceptación final. La identificación de paquetes provisionales anteriores no acredita cierre.'],
        ['Originales oficiales', 'docs/examen-oficial/originales/ y docs/examen-oficial/anexos/'],
        ['Procedencia editorial', 'Edición asistida por Codex. No cambia requisitos ni atribuye firmas o trabajo humano.']])
    heading('1.1 Qué funciona y qué falta', 2)
    paragraph('La aplicación provisional incluye sesiones, ocho cuentas ficticias, formulario de solicitudes, transiciones por rol, reserva transaccional, idempotencia, inventario, listados, tablero e historial. El asistente integrado utiliza únicamente fallback determinista, visible como tal; no se ha demostrado una salida de un LLM real en esta aplicación.')
    paragraph('Pendientes de cierre: revisión humana de las versiones correspondientes, proveedor IA y validación estricta completa, carga exigida, pruebas de todos los criterios, restauración y aceptación manual. Adelantar implementación por instrucción del coordinador no subsana la revisión previa requerida por el examen.')
    source('docs/examen-oficial/anexos/09_pdf-delivery-template.md', sha)
    heading('1.2 Aportes comprobados directamente en GitHub', 2)
    paragraph('Consulta de evidencia: ' + github['fetchedAtUtc'] + '. Se revisaron PR, comentarios, revisiones y ejecuciones; se conserva una instantánea local verificable. Un PR abierto no significa aporte fusionado ni aprobación del paquete. Los nombres de cuenta identifican publicaciones, no prueban autoría manual de contenido asistido por IA.')
    table(['PR / cuenta', 'Artefacto o aporte', 'SHA / estado'], [[f"#{pr['number']} · {pr['author']}", pr['title'], pr['sha'] + (' · borrador' if pr['draft'] else ' · abierto') + (' · fusionado' if pr['merged_at'] else ' · sin fusionar')] for pr in sorted(github['pullRequests'], key=lambda p: p['number'])])
    paragraph('El PR #11 de Vega incorpora el workflow de verificación documental y AI-008. Se incluye su registro con procedencia exacta, sin fusionar la rama ni presentar CI documental como prueba de aplicación. El PR #12 contiene la aplicación y la edición de este informe.')
    for pr in sorted(github['pullRequests'], key=lambda p: p['number']):
        add('link', (f"PR #{pr['number']} · {pr['sha'][:12]}", pr['url']))

    heading('2. Cinco historias de usuario y especificaciones')
    paragraph('Las siguientes especificaciones proceden de los JSON generados a partir de las fuentes Markdown. Se conserva el contenido funcional y el estado declarado. Los 73 criterios son objetivos de aceptación; no deben interpretarse como 73 pruebas superadas.')
    specs = [json.loads(p.read_text(encoding='utf-8')) for p in sorted((ROOT / 'specs/generated').glob('SPEC-*.json'))]
    table(['Especificación', 'Versión / criterios', 'Revisión'], [[s['id'] + ' · ' + s['name'], s['version'] + ' / ' + str(len(s['acceptanceCriteria'])), 'Pendiente según fuente; ver registro humano por SHA'] for s in specs])
    for spec in specs:
        add('break')
        heading(spec['id'] + ' · ' + spec['name'], 2)
        paragraph('Versión: ' + spec['version'] + '. Estado: ' + spec['status'])
        story = spec['story']
        add('callout', f"{story['id']} — Como {story['as']}, quiero {story['want']}, para {story['soThat']}.")
        paragraph(spec['description'])
        paragraph('Requisitos relacionados: ' + ', '.join(spec['requirements']) + '.')
        for key, label in [('actors', 'Actores'), ('scope', 'Alcance'), ('exclusions', 'Fuera del alcance')]:
            heading(label, 3)
            for value in spec[key]:
                add('bullet', value)
        heading('Entradas y validaciones', 3)
        table(['Entrada', 'Tipo', 'Validación'], [[v['field'], v['type'], '\n'.join(v['validation'])] for v in spec['inputs']])
        for key, label in [('process', 'Proceso'), ('businessRules', 'Reglas de negocio'), ('edgeCases', 'Casos límite'), ('assumptions', 'Supuestos registrados'), ('openQuestions', 'Preguntas registradas'), ('dependencies', 'Dependencias')]:
            heading(label, 3)
            values = spec[key]
            if not values:
                paragraph('Sin entradas adicionales en esta versión de la fuente; no acredita resolución humana de otros pendientes.')
            for value in values:
                add('bullet', value if isinstance(value, str) else json.dumps(value, ensure_ascii=False))
        heading('Salidas', 3)
        table(['Condición', 'Resultado'], [[v['condition'], v['result']] for v in spec['outputs']])
        heading('Criterios de aceptación', 3)
        for ac in spec['acceptanceCriteria']:
            add('ac', ac)
        source(f"specs/{spec['id']}.md", sha)

    heading('3. Arquitectura, datos y contratos')
    heading('3.1 Estructura y límites', 2)
    paragraph('La propuesta es un monolito modular Laravel con PostgreSQL y una interfaz Blade/JavaScript del mismo origen. El servidor resuelve usuario, rol y organización; la interfaz solo presenta acciones permitidas. Estado, stock, auditoría y resultado idempotente se confirman en una transacción de negocio. La aprobación del resumen IA no reserva ni despacha kits.')
    table(['Módulo', 'Responsabilidad', 'Límite relevante'], [
        ['Identity', 'Autenticación y sesión', 'Rol/organización proceden del servidor'],
        ['Requests', 'Cabeceras y líneas, edición de borradores', 'Dueño y organización verificables'],
        ['Inventory', 'Reserva, despacho, cancelación y entrega', 'Bloqueos, invariantes y transacción'],
        ['Queries / Audit', 'Consultas y eventos', 'Alcance propio u organizacional según rol'],
        ['AI', 'Resumen alternativo y aprobación', 'Sin escritura privilegiada de stock; proveedor real pendiente']])
    paragraph('Diferencia de implementación que requiere revisión: ADR-001 propone capas internas y routes/api.php; la versión provisional utiliza clases por módulo y rutas en routes/web.php para sesiones y CSRF. No se presenta esta diferencia como una nueva decisión del cliente ni como conformidad completa con el ADR.')
    heading('3.2 Diagramas y modelo de datos', 2)
    table(['Artefacto', 'Qué documenta'], [
        ['contexto.md', 'Actores, sistema y servicios externos'], ['componentes.md', 'Módulos y fronteras'],
        ['datos.md', 'Entidades y relaciones'], ['estados.md', 'Estados y transiciones permitidas'],
        ['secuencia-reserva-concurrente.md', 'Competencia por existencias y resultado esperado'],
        ['secuencia-ia-con-fallo.md', 'Validación, fallo del proveedor y fallback']])
    paragraph('Los seis diagramas editables Mermaid están versionados en docs/architecture/. Las figuras siguientes se renderizan directamente desde esas fuentes, sin cambiar nodos, relaciones ni decisiones. Representan el diseño documentado; no certifican implementación ni aprobación de sus elementos propuestos.')
    for file, caption in [('contexto', 'Contexto y fronteras del sistema'), ('componentes', 'Componentes y dependencias'), ('datos', 'Entidades y relaciones del modelo propuesto'), ('estados', 'Estados y transiciones de solicitudes'), ('secuencia-reserva-concurrente', 'Secuencia de reserva concurrente'), ('secuencia-ia-con-fallo', 'Secuencia de IA con fallo y fallback')]:
        if not (ROOT / 'docs/architecture/figures' / (file + '.png')).exists():
            raise SystemExit('Falta renderizar diagrama: ' + file)
        add('diagram', (ROOT / 'docs/architecture/figures' / (file + '.png'), caption))
    for p in sorted((ROOT / 'docs/architecture').glob('*.md')):
        source(p.relative_to(ROOT).as_posix(), sha)
    table(['Entidad', 'Datos principales', 'Control'], [
        ['organizations / users', 'Organización; usuario, rol y hash', 'Ocho cuentas ficticias iniciales'],
        ['inventory_items', 'Organización + SKU; físico, reservado, activo', 'CHECK: 0 ≤ reservado ≤ físico'],
        ['requests / request_lines', 'UUID, dueño, estado, fecha; SKU y cantidad', 'FK compuestas y cantidades positivas'],
        ['audit_events', 'Actor, recurso, acción, UTC y correlationId', 'Eventos inmutables mediante trigger'],
        ['idempotency_records', 'Usuario/org/operación/clave, hash, respuesta y TTL', 'Ámbito único; reserva ≥24 h'],
        ['ai_summaries / ai_attempts', 'Resumen, origen, estado, aprobación e intento', 'Fallback identificado y aprobación separada']])
    heading('3.3 Contrato HTTP integrado', 2)
    api = yaml.safe_load((ROOT / 'contracts/openapi.yaml').read_text(encoding='utf-8'))
    paragraph('OpenAPI 3.1.1 integrado: 20 rutas y 22 operaciones. La siguiente tabla desarrolla la superficie HTTP; los esquemas y reglas completos se conservan en el contrato enlazado. El metadato introductorio del contrato aún contiene referencias históricas de G2; requiere alineación editorial, no una aprobación implícita.')
    rows = []
    for path, item in api['paths'].items():
        for method, op in item.items():
            if method in ['get', 'post', 'patch', 'put', 'delete']:
                rows.append([method.upper(), path, op.get('summary', op.get('operationId', ''))])
    table(['Método', 'Ruta', 'Propósito'], rows)
    paragraph('Seguridad: cookie de sesión y CSRF en escrituras; 401 sin identidad, 403 cuando el rol o dueño carece de permiso y 404 para recurso inexistente o ajeno a la organización según el caso. La reserva exige Idempotency-Key. Los errores incluyen code, message y correlationId; validación puede añadir field. Solicitudes y listados se filtran en servidor antes de contar y paginar.')
    paragraph('CL-16/17/18: se registra la confirmación afirmativa comunicada y la captura aportada por el coordinador; Coordinador consulta inventario/historial organizacional, Solicitante consulta historial y tablero propios. No se amplía acceso a datos ajenos. La evidencia debe conservarse con su procedencia real.')
    source('contracts/openapi.yaml', sha)
    source('contracts/openapi-ia.yaml', sha)
    source('contracts/ai-output.schema.json', sha)
    heading('3.4 Decisiones de arquitectura', 2)
    paragraph('Se reproducen los cuatro ADR con estilos editoriales y tablas. Su estado propuesto o pendiente se conserva; las expresiones históricas sobre trabajo futuro no equivalen a constatación de la implementación actual.')
    for p in sorted((ROOT / 'docs/adr').glob('ADR-*.md')):
        markdown(p.relative_to(ROOT).as_posix())
        source(p.relative_to(ROOT).as_posix(), sha)

    heading('4. Uso crítico de IA y verificación')
    paragraph('La bitácora diferencia el uso de IA para preparar el proyecto del asistente integrado en la aplicación. Se incluyen ocho registros existentes, siete de la consolidación y AI-008 del PR #11; la autoría y las verificaciones se reproducen como están documentadas, sin crear decisiones humanas nuevas. El asistente actual solo produce fallback: falta demostrar el proveedor real y su validación estricta.')
    for p in sorted((ROOT / 'docs/evidence/ai-log').glob('AI-*.md')):
        heading(p.stem, 2)
        markdown(p.relative_to(ROOT).as_posix(), 3)
    heading('AI-008 · CI documental publicado por Vega', 2)
    paragraph('Fuente adicional recuperada del PR #11, SHA 435651b48e37729133e4db15f9ec8540123c7162. El contenido siguiente se reproduce como registro publicado por su cuenta; no se convierte en revisión humana nueva de este informe.')
    markdown('docs/evidence/github/sources/AI-008-ci-documental.md', 3)
    add('link', ('Fuente original de AI-008 en GitHub', REPO + '/blob/435651b48e37729133e4db15f9ec8540123c7162/docs/evidence/ai-log/AI-008-ci-documental.md'))

    heading('5. Validación, trazabilidad y resultados')
    add('callout', 'Ocho pruebas PHP correctas no significan 73 criterios aceptados. Los resultados se vinculan a sus versiones; no se trasladan automáticamente a una versión posterior.')
    table(['Prueba / versión', 'Esperado', 'Observado y evidencia'], [
        ['PostgreSQL, corrección dd806db', 'Historial propio sin error; sin fuga entre dueños u organizaciones', '8 correctas, 67 aserciones, 1,47 s; phpunit-audit.xml'],
        ['Pint, corrección dd806db', 'Estilo conforme', '35 archivos correctos'],
        ['E2E, candidato f33183a', 'Entrega, teclado, rechazo y CSRF a 390/1366', '4 correctas, 1,6 min; playwright.json'],
        ['Concurrencia HTTP, f33183a', 'Dos reservas de 3 con físico 4: una gana', '200/409; un RESERVE; stock 4/3/1; cancelar restaura 4/0/4'],
        ['CI, f33183a', 'Suite relacional, navegador, concurrencia y documentos', 'Success: runs 38631036099 y 38631043185'],
        ['Documentos y contrato', 'Fuentes coherentes y OpenAPI válido', '14 pruebas documentales; 20 rutas / 22 operaciones'],
        ['LLM real / carga / restauración', 'Evidencia de operación y objetivos del examen', 'Pendiente; no se declara superado']])
    paragraph('Entorno observado: Windows, Docker con Linux, PHP 8.5.0, Laravel fijado en composer.lock, PostgreSQL 17.6, Node 22.22.2 y Playwright 1.56.1. La suite relacional utiliza nexo_test, separada de nexo de demostración. CI se ejecutó en Ubuntu. Los intentos interrumpidos con Edge no se contabilizan como suite correcta; la ejecución completa anterior utilizó Chromium aislado.')
    heading('5.1 CI comprobada en GitHub para la corrección actual', 2)
    paragraph('Consulta: ' + current_ci['checkedAtUtc'] + '. La versión 02ef410 incluye la corrección dd806db y la primera edición Word/PDF. La evidencia remota siguiente amplía los resultados anteriores; no acredita el LLM real, carga ni aceptación de todos los AC.')
    latest = [run for run in current_ci['runs'] if run['head_sha'] == '02ef410ed11d26fa651e405f1563879e987aea97']
    table(['Ejecución', 'SHA', 'Estado observado'], [[str(run['id']), run['head_sha'], run['status'] + ' / ' + str(run['conclusion'])] for run in latest])
    for run in latest:
        add('link', ('Ver ejecución CI ' + str(run['id']), run['html_url']))
    checks = json.loads((ROOT / 'docs/evidence/github/pr-11-checks.json').read_text(encoding='utf-8'))
    paragraph('Comprobación adicional del PR #11 en 435651b: GitHub registra los dos checks siguientes; su alcance es documental y secretos, separado del pipeline de aplicación.')
    table(['Check de Vega', 'Estado', 'Finalizado UTC'], [[check['name'], str(check['conclusion']), str(check['completed_at'])] for check in checks])
    for check in checks:
        add('link', ('Evidencia GitHub: ' + check['name'], check['html_url']))
    heading('5.2 Defecto observado y corrección', 2)
    markdown('docs/evidence/runtime/correccion-auditoria.md', 3)
    heading('5.3 Trazabilidad SPEC frente a código', 2)
    modules = ['Identity / middleware', 'Requests / interfaz', 'Inventory / transacciones', 'Queries / interfaz', 'AI / fallback']
    table(['SPEC', 'Código relacionado', 'Límite de la evidencia'], [[s['id'] + f" ({len(s['acceptanceCriteria'])} AC)", modules[i], 'Escenarios parciales ejecutados; aceptación completa pendiente'] for i, s in enumerate(specs)])
    paragraph('La matriz completa de 73 AC se conserva en docs/evidence/runtime/spec-vs-code.csv, con estado NOT_VERIFIED_IN_FULL. Cada AC completo está descrito en el apartado 2. La matriz documental enlaza pruebas planificadas; solo los reportes ejecutados acreditan resultados observados.')
    source('docs/evidence/runtime/spec-vs-code.csv', sha)
    source('docs/coverage-matrix.md', sha)
    source('docs/evidence/runtime/verification.md', sha)
    for name in ['solicitudes-desktop.png', 'solicitudes-mobile.png']:
        add('image', (ROOT / 'docs/evidence/runtime' / name, 'Captura real anterior a la corrección de auditoría; datos ficticios. ' + ('1366 px.' if 'desktop' in name else '390 px.')))

    heading('6. Riesgos y controles')
    paragraph('La siguiente lectura conserva los riesgos registrados y distingue controles observados de pendientes. Las personas propuestas no se consideran autoras de trabajo no realizado; el riesgo residual requiere evaluación humana.')
    rows = []
    risk_evidence = {
        'R-01': 'Reserva concurrente f33183a: 200/409 y un evento. Residual por evaluar: faltan más intercalados y carga.',
        'R-02': 'Suite relacional con aislamiento; regresión dd806db verifica historial propio y ajeno. Residual por evaluar: matriz completa pendiente.',
        'R-03': 'Replay y conflicto de clave comprobados en PostgreSQL. Residual por evaluar: TTL y escenarios concurrentes adicionales pendientes.',
        'R-04': 'Proveedor y validador completo no implementados; riesgo alto hasta probar salida externa e inyección.',
        'R-05': 'Fallback y aprobación sin stock comprobados. Residual por evaluar: falta fallo real del proveedor y timeout.',
        'R-06': 'CI en GitHub correcto para 02ef410; Docker y lockfiles existentes. Residual por evaluar: aceptación de clon limpio y recuperación pendientes.',
        'R-07': 'Originales oficiales recibidos y cotejo documental existente. Residual: vigilar cambios de versión; no afirmar revisión humana completa.',
        'R-08': 'Implementación adelantada por autorización del coordinador. Residual alto: no sustituye el gate previo exigido.',
        'R-09': 'Snapshot GitHub y reportes separan publicaciones y pruebas de tareas propuestas. Residual: validación humana de la contribución personal pendiente.',
        'R-10': 'Cuatro E2E a 390/1366 pasaron en f33183a y la CI de 02ef410 es correcta. Residual: accesibilidad manual y carga/p95 pendientes.',
        'R-11': 'Transacciones y rechazo parcial comprobados. Residual por evaluar: falta fallo inducido exhaustivo de auditoría antes de commit.',
        'R-12': 'Confirmación afirmativa reportada con captura de CL-16/17/18. No se mantiene la consulta como no enviada; conservar evidencia original y alcance concreto.'
    }
    for line in (ROOT / 'docs/risks/risk-register.md').read_text(encoding='utf-8').splitlines():
        if line.startswith('| R-'):
            c = [v.strip() for v in line.strip('|').split('|')]
            heading(c[0] + ' · ' + c[1], 3)
            paragraph('Probabilidad: ' + c[3] + '. Impacto: ' + c[2] + '. Responsable en el registro: ' + c[5] + '.')
            paragraph('Control documentado: ' + c[4] + '. Evidencia en la fuente histórica: ' + c[6] + '. Residual registrado: ' + c[7] + '.')
            paragraph('Estado actual del control: ' + risk_evidence[c[0]])
    paragraph('Actualización de lectura: hay evidencia parcial de reserva, aislamiento, idempotencia y ejecución; el material oficial ya está recibido y CL-16/17/18 cuentan con respuesta reportada. El registro fuente contiene textos históricos pendientes de armonización. Se mantiene alto el riesgo de proceso por cierre humano pendiente; IA, carga y recuperación siguen sin aceptación completa.')
    source('docs/risks/risk-register.md', sha)

    heading('7. Ejecución reproducible y recuperación')
    markdown('docs/delivery/runbook-provisional.md', 2)
    heading('7.1 Corrección de auditoría en la instancia local', 2)
    paragraph('La instancia local recibió QueryController corregido sin modificar datos ni esquema. Un arranque limpio debe reconstruir la imagen desde el código identificado en esta edición. Verificar historial del Solicitante tras iniciar sesión. Las contraseñas expuestas en el runbook son exclusivamente ficticias de demostración.')

    heading('8. Aceptación y pendientes verificables')
    markdown('docs/delivery/aceptacion-provisional.md', 2)
    markdown('docs/delivery/checklist-entrega.md', 2)
    paragraph('Revisor humano: pendiente. Fecha UTC de decisión: pendiente. SHA aceptado: pendiente. Observaciones y decisión: deben ser publicadas personalmente por el revisor correspondiente. Esta edición no firma por integrantes ni declara revisiones anteriores a la implementación.')

    heading('9. Participación y evidencias del equipo')
    markdown('docs/delivery/participacion-verificada.md', 2)
    paragraph('Aporte adicional comprobado: josevegalago2006-cpu publicó el PR #11 (workflow de CI documental y registro AI-008), aún abierto en borrador. Esta publicación se distingue del pipeline de aplicación del PR #12.')
    heading('9.1 Revisiones y decisiones publicadas, con su alcance', 2)
    paragraph('Las revisiones formales recuperadas tienen estado COMMENTED; no se contabilizan como APPROVED de GitHub. Se transcribe el texto publicado para conservar sus decisiones y observaciones concretas. La revisión asistida del PR #7 no acredita una revisión personal de Conde José. Las aprobaciones textuales de versiones previas no se transfieren al SHA actual.')
    for pr in sorted(github['pullRequests'], key=lambda p: p['number']):
        for review in pr['reviews']:
            heading(f"PR #{pr['number']} · {review['author']} · {review['state']}", 3)
            paragraph('SHA revisado: ' + str(review['sha']) + '. Fecha GitHub: ' + str(review['date']) + '.')
            for line in str(review['body']).splitlines():
                if line.strip(): paragraph(line)
            add('link', ('Revisión original en GitHub', review['url']))
        if pr['number'] == 10:
            for comment in pr['comments']:
                if comment['author'] == 'karollCantillo':
                    heading('Decisión textual de Carolina en PR #10', 3)
                    paragraph('Publicada: ' + str(comment['date']) + '. Alcance declarado: SPEC-04 y plan/tareas en c70e431; no cierre global del candidato actual.')
                    for line in comment['body'].splitlines():
                        if line.strip(): paragraph(line)
                    add('link', ('Comentario original de Carolina', comment['url']))
    paragraph('Los enlaces a PR se consultan en el repositorio: #2/#8 (Edier), #3/#4 (Vega), #7/#10 (Carolina y consolidación), #12 (frontend/integración provisional). Las revisiones válidas requieren autor real, SHA y decisión. Las tareas propuestas no acreditan su realización.')
    source('docs/team/plan-colaboracion.md', sha)
    heading('Fuentes y control de edición', 2)
    paragraph('Esta edición transforma la presentación, sin sustituir las fuentes oficiales ni modificar las cinco SPECS. El Word es editable; si se cambia su contenido, el PDF debe regenerarse y cotejarse para que los cuatro integrantes entreguen el mismo archivo final. La portada no incluye institución, curso o docente no confirmados.')


def clean(text):
    return re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', str(text)).replace('**', '').replace('`', '')


def word(sha):
    d = Document()
    from lxml import etree
    for part in d.part.package.parts:
        if str(part.partname).endswith('/theme/theme1.xml'):
            theme = etree.fromstring(part.blob)
            accent = theme.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}accent1/{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr')
            if accent is not None:
                accent.set('val', GREEN)
                part._blob = etree.tostring(theme, xml_declaration=True, encoding='UTF-8', standalone=True)
    update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); d.settings.element.append(update)
    sec = d.sections[0]
    sec.page_height, sec.page_width = Inches(11.7), Inches(8.3)
    sec.top_margin = sec.bottom_margin = Inches(.75)
    sec.left_margin = sec.right_margin = Inches(.8)
    for name in ['Normal', 'Body Text', 'List Bullet']:
        d.styles[name].font.name = 'Calibri'
        d.styles[name].font.size = Pt(10.5)
        d.styles[name].paragraph_format.space_after = Pt(6)
    d.styles['Normal'].paragraph_format.line_spacing = 1.12
    for n, size in [(1, 19), (2, 14), (3, 11)]:
        st = d.styles[f'Heading {n}']
        st.font.name = 'Calibri'
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(GREEN)
        st.paragraph_format.keep_with_next = True
    sec.header.paragraphs[0].text = 'NEXOENTREGA 360     /     INFORME DE PROYECTO'
    sec.header.paragraphs[0].runs[0].font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.text = 'PROVISIONAL · 14 OCT 2026                                           '
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
    footer.runs[0].font.size = Pt(8)
    def link(label, url):
        p = d.add_paragraph()
        h = OxmlElement('w:hyperlink')
        rel = d.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
        h.set(qn('r:id'), rel)
        r = OxmlElement('w:r'); pr = OxmlElement('w:rPr'); col = OxmlElement('w:color'); col.set(qn('w:val'), '2E6650'); pr.append(col); r.append(pr)
        t = OxmlElement('w:t'); t.text = label; r.append(t); h.append(r); p._p.append(h)
    for kind, value in blocks:
        if kind == 'cover':
            d.add_paragraph('INGENIERÍA DE SOFTWARE · EXAMEN DE EQUIPO', 'Subtitle')
            d.add_paragraph('EXAMEN · SPEC DRIVEN DEVELOPMENT CON IA', 'Subtitle')
            p = d.add_paragraph(); p.paragraph_format.space_before = Pt(45)
            r = p.add_run('NexoEntrega\n360'); r.bold = True; r.font.size = Pt(44); r.font.color.rgb = RGBColor.from_string(GREEN)
            d.add_paragraph('Informe técnico y evidencias', 'Title')
            d.add_paragraph('Especificaciones · Arquitectura · Validación · Equipo', 'Subtitle')
            d.add_paragraph('EDICIÓN PROVISIONAL PARA REVISIÓN — 14 DE OCTUBRE DE 2026')
            t = d.add_table(rows=0, cols=2); t.style = 'Light Shading Accent 1'
            for a, b in [('Cantillo Carolina', '20231153656'), ('Castro Edier', '202311254074'), ('Conde José', '202321256268'), ('Vega José', '202321256198')]:
                c = t.add_row().cells; c[0].text = a; c[1].text = b
            d.add_paragraph('Código de referencia: ' + sha)
            d.add_paragraph('Sin aceptación final ni aprobación humana implícita.')
            d.add_page_break()
        elif kind == 'toc':
            d.add_heading('Índice', 1)
            field = d.add_paragraph()
            begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
            instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' TOC \\o "1-2" \\h \\z \\u '
            separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
            for item in [begin, instr, separate]:
                run = OxmlElement('w:r'); run.append(item); field._p.append(run)
            for k, v in blocks:
                if k == 'h1': d.add_paragraph(v)
            end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end'); run = OxmlElement('w:r'); run.append(end); d.paragraphs[-1]._p.append(run)
            d.add_paragraph('Índice automático de Word. Si no se actualiza al abrir: Ctrl+A y F9, o clic derecho en el índice → Actualizar toda la tabla. Los títulos también aparecen en el panel Navegación.')
            d.add_page_break()
        elif kind.startswith('h'):
            if kind == 'h1' and d.paragraphs[-1].text:
                d.add_page_break()
            d.add_heading(clean(value), int(kind[1]))
        elif kind == 'break': d.add_page_break()
        elif kind == 'table':
            t = d.add_table(rows=0, cols=len(value[0])); t.style = 'Light Shading Accent 1'
            for i, row in enumerate(value):
                cells = t.add_row().cells
                if i == 0:
                    repeat = OxmlElement('w:tblHeader'); t.rows[-1]._tr.get_or_add_trPr().append(repeat)
                for c, text in zip(cells, row):
                    c.text = clean(text)
                    for p in c.paragraphs:
                        for r in p.runs:
                            r.font.size = Pt(9)
                            if i == 0: r.bold = True
            d.add_paragraph()
        elif kind == 'ac':
            p = d.add_paragraph(value['id'] + ' · ' + value['kind'] + ' · ' + ', '.join(value['requirements']), 'Heading 3')
            for key, label in [('given', 'Dado'), ('when', 'Cuando'), ('then', 'Entonces')]:
                p = d.add_paragraph(); p.add_run(label + ': ').bold = True; p.add_run(value[key])
        elif kind == 'link': link(*value)
        elif kind in ['image', 'diagram']:
            p, caption = value
            if p.exists():
                if kind == 'diagram':
                    from PIL import Image as PILImage
                    if not d.paragraphs[-1]._p.xpath('.//w:br[@w:type="page"]'):
                        d.add_page_break()
                    d.add_heading(caption, 3)
                    with PILImage.open(p) as image:
                        width, height = image.size
                    scale = min(6.3 / width, 7.6 / height)
                    d.add_picture(str(p), width=Inches(width * scale), height=Inches(height * scale))
                else:
                    d.add_picture(str(p), width=Inches(2.55 if 'mobile' in p.name else 6.3))
                d.add_paragraph(caption, 'Caption')
                if kind == 'diagram': d.add_page_break()
        elif kind == 'bullet': d.add_paragraph(clean(value), 'List Bullet')
        elif kind == 'code':
            p = d.add_paragraph(clean(value)); p.paragraph_format.space_after = Pt(2)
            for r in p.runs: r.font.name = 'Consolas'; r.font.size = Pt(8)
        else:
            p = d.add_paragraph(clean(value))
            if kind == 'callout':
                p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(12)
                for r in p.runs: r.bold = True; r.font.color.rgb = RGBColor.from_string(GREEN)
    d.core_properties.title = 'NexoEntrega 360 — Informe técnico y evidencias'
    d.core_properties.author = 'Equipo NexoEntrega 360 · edición asistida por Codex'
    d.save(OUT / 'Examen_SDD_IA_Equipo_NexoEntrega360.docx')


class Report(BaseDocTemplate):
    def afterFlowable(self, flowable):
        if isinstance(flowable, Paragraph) and flowable.style.name in ['Section', 'Subsection']:
            level = 0 if flowable.style.name == 'Section' else 1
            key = 'heading-' + str(self.seq.nextf('heading'))
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(flowable.getPlainText(), key, level, False)
            self.notify('TOCEntry', (level, flowable.getPlainText(), self.page, key))


def pdf(sha):
    for name, file in [('Nexo', 'arial.ttf'), ('NexoBold', 'arialbd.ttf'), ('NexoItalic', 'ariali.ttf')]:
        pdfmetrics.registerFont(TTFont(name, 'C:/Windows/Fonts/' + file))
    pdfmetrics.registerFontFamily('Nexo', normal='Nexo', bold='NexoBold', italic='NexoItalic', boldItalic='NexoBold')
    st = getSampleStyleSheet()
    for s in st.byName.values(): s.fontName = 'Nexo'
    st.add(ParagraphStyle('Text', fontName='Nexo', fontSize=9.5, leading=14, spaceAfter=7, textColor=colors.HexColor('#263B35')))
    st.add(ParagraphStyle('Section', parent=st['Heading1'], fontName='NexoBold', fontSize=20, leading=25, spaceAfter=16, textColor=colors.HexColor('#' + GREEN)))
    st.add(ParagraphStyle('Subsection', parent=st['Heading2'], fontName='NexoBold', fontSize=13, leading=18, spaceBefore=12, spaceAfter=8, textColor=colors.HexColor('#' + GREEN)))
    st.add(ParagraphStyle('Minor', parent=st['Heading3'], fontName='NexoBold', fontSize=10, leading=14, spaceBefore=9, textColor=colors.HexColor('#' + GREEN)))
    st.add(ParagraphStyle('Cell', parent=st['Text'], fontSize=8, leading=11, spaceAfter=0, wordWrap='CJK'))
    st.add(ParagraphStyle('Callout', parent=st['Text'], fontName='NexoBold', backColor=colors.HexColor('#EDF3E9'), borderPadding=10, spaceBefore=9, spaceAfter=15))
    st.add(ParagraphStyle('CodeBlock', parent=st['Text'], fontName='Courier', fontSize=7.2, leading=10, backColor=colors.HexColor('#F3F5F3'), wordWrap='CJK'))
    def p(text, style='Text'):
        if text in ['Informe técnico y evidencias', 'Índice']:
            style = 'Minor'
        return Paragraph(escape(clean(text)).replace('\n', '<br/>'), st[style])
    parts = []
    for kind, value in blocks:
        if kind == 'cover':
            parts.append(p('INGENIERÍA DE SOFTWARE · EXAMEN DE EQUIPO'))
            parts += [Spacer(1, 35), p('EXAMEN · SDD CON IA'), Spacer(1, 35), Paragraph('NexoEntrega<br/>360', ParagraphStyle('Cover', fontName='NexoBold', fontSize=49, leading=56, textColor=colors.HexColor('#' + GREEN))), Spacer(1, 20), p('Informe técnico y evidencias', 'Subsection'), p('Especificaciones · Arquitectura · Validación · Equipo'), Spacer(1, 26), p('EDICIÓN PROVISIONAL PARA REVISIÓN', 'Callout'), p('14 de octubre de 2026'), Spacer(1, 18)]
            for name, code in [('Cantillo Carolina', '20231153656'), ('Castro Edier', '202311254074'), ('Conde José', '202321256268'), ('Vega José', '202321256198')]:
                parts.append(p(name + '  ·  ' + code))
            parts += [Spacer(1, 20), p('Código de referencia'), p(sha, 'Cell'), p('Sin aceptación final ni aprobación humana implícita.'), PageBreak()]
        elif kind == 'toc':
            parts += [p('Índice', 'Subsection')]
            toc = TableOfContents(); toc.levelStyles = [ParagraphStyle('Toc0', fontName='NexoBold', fontSize=10, leading=16, spaceBefore=9), ParagraphStyle('Toc1', fontName='Nexo', fontSize=8.5, leading=13, leftIndent=12)]
            parts += [toc, PageBreak()]
        elif kind == 'h1':
            if parts and not isinstance(parts[-1], PageBreak): parts.append(PageBreak())
            parts.append(p(value, 'Section'))
        elif kind in ['h2', 'h3']: parts.append(p(value, 'Subsection' if kind == 'h2' else 'Minor'))
        elif kind == 'break': parts.append(PageBreak())
        elif kind == 'table':
            n = len(value[0]); widths = [495 / n] * n
            if n == 2: widths = [155, 340]
            if n == 3: widths = [115, 150, 230]
            if n == 3 and value[0][0] == 'Método': widths = [48, 235, 212]
            data = [[p(c, 'Cell') for c in row] for row in value]
            t = Table(data, colWidths=widths, repeatRows=1, splitByRow=1, splitInRow=1, hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DDEBDD')), ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F7F3')]), ('LEFTPADDING', (0, 0), (-1, -1), 8), ('RIGHTPADDING', (0, 0), (-1, -1), 8), ('TOPPADDING', (0, 0), (-1, -1), 7), ('BOTTOMPADDING', (0, 0), (-1, -1), 7), ('LINEBELOW', (0, 0), (-1, 0), .7, colors.HexColor('#2E6650'))]))
            parts += [t, Spacer(1, 10)]
        elif kind == 'ac':
            parts.append(KeepTogether([p(value['id'] + ' · ' + value['kind'] + ' · ' + ', '.join(value['requirements']), 'Minor'), p('Dado: ' + value['given']), p('Cuando: ' + value['when']), p('Entonces: ' + value['then'])]))
        elif kind == 'link':
            label, url = value
            parts.append(Paragraph('<link href="' + escape(url, quote=True) + '" color="#2E6650">Fuente: ' + escape(label) + '</link>', st['Cell']))
            parts.append(Spacer(1, 7))
        elif kind in ['image', 'diagram']:
            file, caption = value
            if file.exists():
                if kind == 'diagram':
                    if not parts or not isinstance(parts[-1], PageBreak): parts.append(PageBreak())
                    parts.append(p(caption, 'Minor'))
                img = Image(str(file)); ratio = min(490 / img.imageWidth, (635 if kind == 'diagram' else 450) / img.imageHeight); img.drawWidth *= ratio; img.drawHeight *= ratio
                parts += [Spacer(1, 12), img, p(caption, 'Cell'), Spacer(1, 12)]
                if kind == 'diagram': parts.append(PageBreak())
        else:
            parts.append(p(('• ' if kind == 'bullet' else '') + str(value), {'callout': 'Callout', 'code': 'CodeBlock'}.get(kind, 'Text')))
    def decorate(canvas, doc):
        canvas.saveState(); canvas.setStrokeColor(colors.HexColor('#DDE7DA')); canvas.line(50, 801, 545, 801)
        canvas.setFont('Nexo', 7.5); canvas.setFillColor(colors.HexColor('#52685D'))
        canvas.drawString(50, 812, 'NEXOENTREGA 360  /  INFORME TÉCNICO')
        canvas.drawString(50, 28, 'PROVISIONAL · 14 OCT 2026 · ' + sha[:12]); canvas.drawRightString(545, 28, str(doc.page)); canvas.restoreState()
    doc = Report(str(OUT / 'Examen_SDD_IA_Equipo_NexoEntrega360.pdf'), pagesize=A4, leftMargin=50, rightMargin=50, topMargin=56, bottomMargin=48, title='NexoEntrega 360 — Informe técnico y evidencias', author='Equipo NexoEntrega 360 · edición asistida por Codex')
    frame = Frame(50, 48, 495, 738, leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)
    doc.addPageTemplates(PageTemplate(id='report', frames=frame, onPage=decorate))
    doc.multiBuild(parts)


def main():
    parser = argparse.ArgumentParser(description=__doc__); parser.add_argument('--sha', required=True); args = parser.parse_args()
    sha = subprocess.check_output(['git', 'rev-parse', args.sha], cwd=ROOT, text=True).strip()
    OUT.mkdir(parents=True, exist_ok=True)
    content(sha); word(sha); pdf(sha)
    manifest = {'status': 'PROVISIONAL_NOT_ACCEPTED', 'sourceSha': sha, 'files': {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in OUT.glob('Examen*')}, 'note': 'Original provisional preserved; Word and PDF generated from same editorial content. No human approval recorded.'}
    (OUT / 'manifest.json').write_text(json.dumps(manifest, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(manifest, indent=2))


if __name__ == '__main__':
    main()
